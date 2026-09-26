from types import SimpleNamespace

from docx import Document

from tools import knowledge_tools


def test_chunk_text_creates_overlapping_chunks():
    text = "A" * 1800

    chunks = knowledge_tools._chunk_text(
        text
    )

    assert len(chunks) >= 2
    assert all(chunks)


def test_list_knowledge_documents_supports_v2_formats(
    monkeypatch,
    tmp_path,
):
    root = tmp_path / "knowledge"
    root.mkdir()

    (root / "README.md").write_text(
        "ignored",
        encoding="utf-8",
    )
    (root / "notes.md").write_text(
        "hello",
        encoding="utf-8",
    )
    (root / "notes.txt").write_text(
        "world",
        encoding="utf-8",
    )
    (root / "guide.pdf").write_bytes(
        b"placeholder"
    )
    (root / "project.docx").write_bytes(
        b"placeholder"
    )
    (root / "binary.bin").write_bytes(
        b"123"
    )

    monkeypatch.setattr(
        knowledge_tools,
        "KNOWLEDGE_ROOT",
        root,
    )

    result = (
        knowledge_tools
        .list_knowledge_documents()
    )

    assert "notes.md" in result
    assert "notes.txt" in result
    assert "guide.pdf" in result
    assert "project.docx" in result
    assert "README.md" not in result
    assert "binary.bin" not in result


def test_read_docx_extracts_paragraphs_and_tables(
    tmp_path,
):
    path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph(
        "Terraform troubleshooting"
    )

    table = document.add_table(
        rows=1,
        cols=2,
    )
    table.cell(
        0,
        0,
    ).text = "Issue"
    table.cell(
        0,
        1,
    ).text = "Resolution"

    document.save(path)

    extracted = (
        knowledge_tools
        ._read_docx(path)
    )

    assert (
        "Terraform troubleshooting"
        in extracted
    )
    assert "Issue | Resolution" in extracted


def test_read_pdf_preserves_page_markers(
    monkeypatch,
    tmp_path,
):
    path = tmp_path / "sample.pdf"
    path.write_bytes(b"pdf")

    fake_reader = SimpleNamespace(
        is_encrypted=False,
        pages=[
            SimpleNamespace(
                extract_text=lambda: "Page one"
            ),
            SimpleNamespace(
                extract_text=lambda: "Page two"
            ),
        ],
    )

    monkeypatch.setattr(
        knowledge_tools,
        "PdfReader",
        lambda _: fake_reader,
    )

    extracted = (
        knowledge_tools
        ._read_pdf(path)
    )

    assert "[Page 1]" in extracted
    assert "Page one" in extracted
    assert "[Page 2]" in extracted
    assert "Page two" in extracted


def test_search_knowledge_reports_missing_index(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        knowledge_tools,
        "INDEX_PATH",
        tmp_path / "missing.sqlite3",
    )

    result = knowledge_tools.search_knowledge(
        "terraform"
    )

    assert "has not been built" in result


def test_knowledge_files_ignore_symlink_escape(
    monkeypatch,
    tmp_path,
):
    root = tmp_path / "knowledge"
    root.mkdir()

    outside = tmp_path / "outside.txt"
    outside.write_text(
        "private",
        encoding="utf-8",
    )

    link = root / "linked.txt"
    link.symlink_to(outside)

    monkeypatch.setattr(
        knowledge_tools,
        "KNOWLEDGE_ROOT",
        root,
    )

    assert (
        link
        not in knowledge_tools._knowledge_files()
    )
