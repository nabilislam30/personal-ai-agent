import hashlib
import json
import math
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from ollama import embed
from pypdf import PdfReader

from config import (
    EMBEDDING_MODEL,
    KNOWLEDGE_CHUNK_OVERLAP,
    KNOWLEDGE_CHUNK_SIZE,
    KNOWLEDGE_EMBED_BATCH_SIZE,
    KNOWLEDGE_INDEX_PATH,
    KNOWLEDGE_MAX_FILE_BYTES,
    KNOWLEDGE_MAX_RESULT_CHARS,
    KNOWLEDGE_MAX_TOP_K,
    KNOWLEDGE_ROOT,
    PROJECT_ROOT,
)


INDEX_PATH = KNOWLEDGE_INDEX_PATH

ALLOWED_EXTENSIONS = {
    ".md",
    ".txt",
    ".pdf",
    ".docx",
}

EXCLUDED_FILENAMES = {
    "README.md",
}

CHUNK_SIZE = KNOWLEDGE_CHUNK_SIZE
CHUNK_OVERLAP = min(
    KNOWLEDGE_CHUNK_OVERLAP,
    max(0, CHUNK_SIZE - 1),
)
EMBED_BATCH_SIZE = KNOWLEDGE_EMBED_BATCH_SIZE
MAX_TOP_K = KNOWLEDGE_MAX_TOP_K
MAX_RESULT_CHARS = KNOWLEDGE_MAX_RESULT_CHARS
MAX_FILE_BYTES = KNOWLEDGE_MAX_FILE_BYTES


def _knowledge_files() -> list[Path]:
    """
    Return supported knowledge files in a deterministic order.

    Symlinks and paths that resolve outside knowledge/ are ignored.
    """

    if not KNOWLEDGE_ROOT.exists():
        return []

    root = KNOWLEDGE_ROOT.resolve()
    files = []

    for path in KNOWLEDGE_ROOT.rglob("*"):
        if not path.is_file():
            continue

        if path.is_symlink():
            continue

        if path.name in EXCLUDED_FILENAMES:
            continue

        if path.name.startswith("."):
            continue

        if path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue

        try:
            resolved = path.resolve()
            resolved.relative_to(root)
        except (OSError, ValueError):
            continue

        try:
            if path.stat().st_size > MAX_FILE_BYTES:
                continue
        except OSError:
            continue

        files.append(path)

    return sorted(
        files,
        key=lambda item: str(
            item.relative_to(KNOWLEDGE_ROOT)
        ).lower(),
    )


def _read_pdf(path: Path) -> str:
    """
    Extract text from a PDF without OCR.

    Scanned/image-only PDFs may return little or no text.
    """

    reader = PdfReader(str(path))

    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as error:
            raise ValueError(
                "Encrypted PDF could not be opened."
            ) from error

        if unlocked == 0:
            raise ValueError(
                "Encrypted PDF requires a password."
            )

    sections = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        try:
            text = page.extract_text() or ""
        except Exception as error:
            raise ValueError(
                f"Could not extract PDF page {page_number}."
            ) from error

        text = text.strip()

        if text:
            sections.append(
                f"[Page {page_number}]\n{text}"
            )

    return "\n\n".join(sections)


def _read_docx(path: Path) -> str:
    """
    Extract paragraph and table text from a DOCX file.
    """

    document = Document(str(path))
    sections = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    for table_number, table in enumerate(
        document.tables,
        start=1,
    ):
        table_rows = []

        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                table_rows.append(
                    " | ".join(cells)
                )

        if table_rows:
            sections.append(
                (
                    f"[Table {table_number}]\n"
                    + "\n".join(table_rows)
                )
            )

    return "\n\n".join(sections)


def _read_knowledge_document(
    path: Path,
) -> str:
    """
    Extract searchable text from a supported knowledge document.
    """

    extension = path.suffix.lower()

    if extension in {
        ".md",
        ".txt",
    }:
        return path.read_text(
            encoding="utf-8"
        )

    if extension == ".pdf":
        return _read_pdf(path)

    if extension == ".docx":
        return _read_docx(path)

    raise ValueError(
        f"Unsupported knowledge file type: {extension}"
    )


def _source_manifest(
    files: list[Path],
) -> str:
    """
    Build a lightweight source-state fingerprint.
    """

    entries = []

    for path in files:
        try:
            stat = path.stat()
        except OSError:
            continue

        entries.append(
            (
                str(
                    path.relative_to(
                        KNOWLEDGE_ROOT
                    )
                ),
                stat.st_size,
                stat.st_mtime_ns,
            )
        )

    payload = json.dumps(
        entries,
        sort_keys=True,
        separators=(",", ":"),
    )

    return hashlib.sha256(
        payload.encode("utf-8")
    ).hexdigest()


def _chunk_text(
    text: str,
) -> list[str]:
    """
    Split text into overlapping chunks while preferring paragraph breaks.
    """

    cleaned = text.strip()

    if not cleaned:
        return []

    chunks = []
    start = 0
    text_length = len(cleaned)

    while start < text_length:
        hard_end = min(
            start + CHUNK_SIZE,
            text_length,
        )

        end = hard_end

        if hard_end < text_length:
            search_start = start + int(
                CHUNK_SIZE * 0.6
            )

            paragraph_break = cleaned.rfind(
                "\n\n",
                search_start,
                hard_end,
            )

            if paragraph_break != -1:
                end = paragraph_break
            else:
                line_break = cleaned.rfind(
                    "\n",
                    search_start,
                    hard_end,
                )

                if line_break != -1:
                    end = line_break

        chunk = cleaned[
            start:end
        ].strip()

        if chunk:
            chunks.append(chunk)

        if end >= text_length:
            break

        next_start = max(
            end - CHUNK_OVERLAP,
            start + 1,
        )

        start = next_start

    return chunks


def _connect() -> sqlite3.Connection:
    """
    Open the local SQLite knowledge index.
    """

    INDEX_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    connection = sqlite3.connect(
        INDEX_PATH
    )

    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_path TEXT NOT NULL,
            chunk_index INTEGER NOT NULL,
            content TEXT NOT NULL,
            content_hash TEXT NOT NULL,
            embedding_json TEXT NOT NULL,
            embedding_model TEXT NOT NULL
        )
        """
    )

    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS metadata (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        )
        """
    )

    connection.commit()

    return connection


def _cosine_similarity(
    left: list[float],
    right: list[float],
) -> float:
    """
    Calculate cosine similarity between two vectors.
    """

    if len(left) != len(right):
        return 0.0

    dot_product = sum(
        a * b
        for a, b in zip(
            left,
            right,
        )
    )

    left_norm = math.sqrt(
        sum(
            value * value
            for value in left
        )
    )

    right_norm = math.sqrt(
        sum(
            value * value
            for value in right
        )
    )

    if (
        left_norm == 0
        or right_norm == 0
    ):
        return 0.0

    return (
        dot_product
        / (left_norm * right_norm)
    )


def list_knowledge_documents() -> str:
    """
    List documents currently available to the local knowledge base.
    """

    files = _knowledge_files()

    if not files:
        return (
            "No knowledge documents were found. "
            "Add .md, .txt, .pdf, or .docx files inside knowledge/."
        )

    lines = [
        "Knowledge documents:",
    ]

    for path in files:
        relative_path = path.relative_to(
            KNOWLEDGE_ROOT
        )

        lines.append(
            f"- {relative_path} "
            f"({path.suffix.lower()}, "
            f"{path.stat().st_size} bytes)"
        )

    return "\n".join(lines)


def knowledge_status() -> str:
    """
    Show source and index status for the local knowledge base.
    """

    files = _knowledge_files()
    current_manifest = _source_manifest(
        files
    )

    if not INDEX_PATH.exists():
        return (
            "Knowledge base status:\n"
            f"- Source documents: {len(files)}\n"
            "- Supported formats: .md, .txt, .pdf, .docx\n"
            "- Index: not built\n"
            "- Indexed chunks: 0\n"
            f"- Embedding model: {EMBEDDING_MODEL}"
        )

    try:
        connection = sqlite3.connect(
            INDEX_PATH
        )

        chunk_count = connection.execute(
            "SELECT COUNT(*) FROM chunks"
        ).fetchone()[0]

        indexed_documents = (
            connection.execute(
                """
                SELECT COUNT(DISTINCT source_path)
                FROM chunks
                """
            ).fetchone()[0]
        )

        metadata_rows = connection.execute(
            "SELECT key, value FROM metadata"
        ).fetchall()

        connection.close()

    except sqlite3.Error as error:
        return (
            "Error reading knowledge index: "
            f"{error}"
        )

    metadata = dict(
        metadata_rows
    )

    indexed_at = metadata.get(
        "indexed_at",
        "unknown",
    )

    indexed_model = metadata.get(
        "embedding_model",
        EMBEDDING_MODEL,
    )

    indexed_manifest = metadata.get(
        "source_manifest",
        "",
    )

    source_state = (
        "up to date"
        if indexed_manifest == current_manifest
        else "changed since last index"
    )

    return (
        "Knowledge base status:\n"
        f"- Source documents: {len(files)}\n"
        f"- Indexed documents: {indexed_documents}\n"
        f"- Indexed chunks: {chunk_count}\n"
        f"- Source state: {source_state}\n"
        f"- Embedding model: {indexed_model}\n"
        f"- Last indexed: {indexed_at}\n"
        f"- Index path: "
        f"{INDEX_PATH.relative_to(PROJECT_ROOT)}"
    )


def index_knowledge() -> str:
    """
    Rebuild the local vector index from knowledge/ documents.

    This writes derived index data only under workspace/.
    Source documents are never modified.
    """

    files = _knowledge_files()

    if not files:
        return (
            "No knowledge documents were found. "
            "Add .md, .txt, .pdf, or .docx files "
            "inside knowledge/ before indexing."
        )

    chunk_records = []
    skipped = []
    indexed_paths = set()

    for path in files:
        relative_path = str(
            path.relative_to(
                KNOWLEDGE_ROOT
            )
        )

        try:
            text = _read_knowledge_document(
                path
            )
        except (
            UnicodeDecodeError,
            OSError,
            ValueError,
        ) as error:
            skipped.append(
                f"{relative_path}: {error}"
            )
            continue

        if not text.strip():
            skipped.append(
                f"{relative_path}: no extractable text"
            )
            continue

        chunks = _chunk_text(
            text
        )

        if not chunks:
            skipped.append(
                f"{relative_path}: no searchable chunks"
            )
            continue

        indexed_paths.add(
            relative_path
        )

        for chunk_index, content in enumerate(
            chunks,
            start=1,
        ):
            content_hash = hashlib.sha256(
                content.encode("utf-8")
            ).hexdigest()

            chunk_records.append(
                {
                    "source_path": relative_path,
                    "chunk_index": chunk_index,
                    "content": content,
                    "content_hash": content_hash,
                }
            )

    if not chunk_records:
        detail = ""

        if skipped:
            detail = (
                "\nSkipped:\n- "
                + "\n- ".join(skipped)
            )

        return (
            "Knowledge documents were found, but no readable "
            "text chunks could be created."
            + detail
        )

    embeddings = []

    try:
        for start in range(
            0,
            len(chunk_records),
            EMBED_BATCH_SIZE,
        ):
            batch = chunk_records[
                start:start
                + EMBED_BATCH_SIZE
            ]

            response = embed(
                model=EMBEDDING_MODEL,
                input=[
                    record["content"]
                    for record in batch
                ],
                truncate=True,
            )

            embeddings.extend(
                [
                    list(vector)
                    for vector in response.embeddings
                ]
            )

    except Exception as error:
        return (
            "Error generating embeddings. "
            f"Ensure Ollama is running and '{EMBEDDING_MODEL}' "
            "is installed. "
            f"Details: {error}"
        )

    if (
        len(embeddings)
        != len(chunk_records)
    ):
        return (
            "Error: embedding count did not match "
            "the number of knowledge chunks."
        )

    try:
        connection = _connect()

        with connection:
            connection.execute(
                "DELETE FROM chunks"
            )

            connection.execute(
                "DELETE FROM metadata"
            )

            for record, vector in zip(
                chunk_records,
                embeddings,
            ):
                connection.execute(
                    """
                    INSERT INTO chunks (
                        source_path,
                        chunk_index,
                        content,
                        content_hash,
                        embedding_json,
                        embedding_model
                    )
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (
                        record["source_path"],
                        record["chunk_index"],
                        record["content"],
                        record["content_hash"],
                        json.dumps(vector),
                        EMBEDDING_MODEL,
                    ),
                )

            metadata = {
                "indexed_at": datetime.now(
                    timezone.utc
                ).isoformat(),
                "embedding_model": EMBEDDING_MODEL,
                "source_manifest": _source_manifest(
                    files
                ),
            }

            for key, value in metadata.items():
                connection.execute(
                    """
                    INSERT INTO metadata (
                        key,
                        value
                    )
                    VALUES (?, ?)
                    """,
                    (
                        key,
                        value,
                    ),
                )

        connection.close()

    except sqlite3.Error as error:
        return (
            "Error writing knowledge index: "
            f"{error}"
        )

    lines = [
        "Knowledge index rebuilt successfully.",
        (
            f"- Documents indexed: "
            f"{len(indexed_paths)}"
        ),
        (
            f"- Chunks indexed: "
            f"{len(chunk_records)}"
        ),
        f"- Embedding model: {EMBEDDING_MODEL}",
        (
            f"- Index: "
            f"{INDEX_PATH.relative_to(PROJECT_ROOT)}"
        ),
    ]

    if skipped:
        lines.append(
            f"- Documents skipped: {len(skipped)}"
        )
        lines.append(
            "Skipped:"
        )
        lines.extend(
            f"- {item}"
            for item in skipped
        )

    return "\n".join(lines)


def search_knowledge(
    query: str,
    top_k: int = 5,
) -> str:
    """
    Search the local vector knowledge base semantically.
    """

    if not query.strip():
        return (
            "Error: Knowledge search query cannot be empty."
        )

    top_k = max(
        1,
        min(
            int(top_k),
            MAX_TOP_K,
        ),
    )

    if not INDEX_PATH.exists():
        return (
            "Knowledge index has not been built yet. "
            "Run index_knowledge after adding "
            "documents to knowledge/."
        )

    try:
        connection = sqlite3.connect(
            INDEX_PATH
        )

        rows = connection.execute(
            """
            SELECT
                source_path,
                chunk_index,
                content,
                embedding_json,
                embedding_model
            FROM chunks
            """
        ).fetchall()

        metadata_rows = connection.execute(
            "SELECT key, value FROM metadata"
        ).fetchall()

        connection.close()

    except sqlite3.Error as error:
        return (
            "Error reading knowledge index: "
            f"{error}"
        )

    if not rows:
        return (
            "Knowledge index contains no chunks. "
            "Run index_knowledge to rebuild it."
        )

    indexed_models = {
        row[4]
        for row in rows
    }

    if indexed_models != {
        EMBEDDING_MODEL
    }:
        return (
            "Knowledge index uses a different embedding model. "
            "Run index_knowledge to rebuild it."
        )

    try:
        response = embed(
            model=EMBEDDING_MODEL,
            input=query,
            truncate=True,
        )

        query_vector = list(
            response.embeddings[0]
        )

    except Exception as error:
        return (
            "Error embedding the knowledge query. "
            f"Ensure Ollama is running and '{EMBEDDING_MODEL}' "
            "is installed. "
            f"Details: {error}"
        )

    ranked = []

    for (
        source_path,
        chunk_index,
        content,
        embedding_json,
        _,
    ) in rows:
        try:
            vector = json.loads(
                embedding_json
            )
        except json.JSONDecodeError:
            continue

        score = _cosine_similarity(
            query_vector,
            vector,
        )

        ranked.append(
            (
                score,
                source_path,
                chunk_index,
                content,
            )
        )

    ranked.sort(
        key=lambda item: item[0],
        reverse=True,
    )

    selected = ranked[
        :top_k
    ]

    if not selected:
        return (
            "No searchable knowledge chunks were available."
        )

    metadata = dict(
        metadata_rows
    )

    current_manifest = _source_manifest(
        _knowledge_files()
    )
    indexed_manifest = metadata.get(
        "source_manifest",
        "",
    )

    output = [
        (
            "Knowledge search results "
            f"for: {query}"
        ),
    ]

    if (
        indexed_manifest
        and indexed_manifest
        != current_manifest
    ):
        output.append(
            (
                "\nWarning: knowledge source files changed "
                "since the index was built. Rebuild the "
                "knowledge index for current results."
            )
        )

    for rank, (
        score,
        source_path,
        chunk_index,
        content,
    ) in enumerate(
        selected,
        start=1,
    ):
        excerpt = content[
            :MAX_RESULT_CHARS
        ]

        if (
            len(content)
            > MAX_RESULT_CHARS
        ):
            excerpt += (
                "\n[Excerpt truncated]"
            )

        output.append(
            (
                f"\n[{rank}] Source: {source_path} "
                f"| Chunk: {chunk_index} "
                f"| Similarity: {score:.4f}\n"
                f"{excerpt}"
            )
        )

    return "\n".join(output)
